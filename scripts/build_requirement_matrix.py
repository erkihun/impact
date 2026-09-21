"""Build a complete SRS requirement inventory without silently dropping IDs."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Impact Consulting Organization Website SRS.docx"
OUTPUT = ROOT / "docs" / "implementation" / "requirements-traceability.md"
REQUIREMENT_ID = re.compile(r"\b(?:FR|DR|IR|NFR|OPS|DOC|TRN)-[A-Z0-9-]+\b")

EVIDENCE: tuple[tuple[tuple[str, ...], str, str], ...] = (
    (
        ("FR-RBA", "FR-ADM", "NFR-SEC"),
        "`app/Actions/Identity/`, `app/Http/Middleware/`, `app/Policies/`, `app/Models/AuditEvent.php`",
        "`IdentityAdministrationTest`, `MfaEnforcementTest`, `SecurityHeadersTest`",
    ),
    (
        ("FR-WFL", "FR-CMS"),
        "`app/Actions/Content/`, `app/Actions/Workflow/`, `app/Jobs/Content/`",
        "`ContentManagementTest`, `ContentWorkflowTest`",
    ),
    (
        ("FR-CON", "FR-LEA"),
        "`app/Actions/Engagement/`, `app/Http/Controllers/Public/EngagementSubmissionController.php`",
        "`EngagementSubmissionTest`, `EngagementOperationsTest`",
    ),
    (
        ("FR-EVT", "FR-CAR"),
        "`app/Actions/Events/`, `app/Actions/Recruitment/`, `app/Http/Controllers/Admin/ApplicationController.php`",
        "`EventRegistrationTest`, `ApplicationSubmissionTest`, `ApplicationOperationsTest`",
    ),
    (
        ("FR-MED", "IR-STO"),
        "`app/Actions/Media/`, `app/Jobs/Media/`, `app/Services/Media/`",
        "`MediaLifecycleTest`",
    ),
    (
        ("FR-SCH", "FR-SEO"),
        "`app/Services/Search/`, `app/Jobs/Search/`, `app/Http/Controllers/Public/SitemapController.php`",
        "`SearchPrivacyAndIndexingTest`, `SeoTest`",
    ),
    (
        ("FR-LNG", "IR-UI"),
        "`app/Http/Middleware/SetLocale.php`, `lang/`, `resources/views/public/`",
        "`LocalizationTest`, browser viewport evidence",
    ),
    (
        (
            "FR-HOM",
            "FR-GEN",
            "FR-ABT",
            "FR-SRV",
            "FR-IND",
            "FR-EXP",
            "FR-CAS",
            "FR-INS",
        ),
        "`app/Http/Controllers/Public/`, `resources/views/public/`, public read models",
        "`LocalizationTest`, `SeoTest`, browser viewport evidence",
    ),
    (
        ("FR-ANA", "FR-NEW", "FR-NTF", "IR-ANL", "IR-EML"),
        "`app/Models/ConsentRecord.php`, `app/Notifications/`, `app/Actions/Search/`",
        "`QueuedAcknowledgementsTest`, public submission tests",
    ),
    (
        ("DR-DQ", "DR-RET", "NFR-PRV", "NFR-AVL", "NFR-MNT", "NFR-USA", "NFR-SCL"),
        "`database/migrations/`, typed Actions and Models, scheduled commands",
        "`tests/Feature/`, `tests/Unit/`",
    ),
    (
        ("NFR-A11Y", "NFR-CMP", "NFR-PER"),
        "`resources/views/`, `resources/css/`, `deploy/`",
        "responsive browser evidence; formal certification pending",
    ),
    (
        ("DOC-",),
        "`README.md`, `docs/`",
        "documentation inspection",
    ),
)


def clean(value: str) -> str:
    return " ".join(value.split()).replace("|", "\\|")


def evidence_for(identifier: str) -> tuple[str, str, str]:
    for prefixes, implementation, tests in EVIDENCE:
        if identifier.startswith(prefixes):
            return "Partial", implementation, tests

    return "Planned", "Backlog; see `module-plan.md`", "Not yet implemented"


def main() -> None:
    document = Document(SOURCE)
    requirements: dict[str, str] = {}

    for table in document.tables:
        for row in table.rows:
            cells = [clean(cell.text) for cell in row.cells]
            joined = " ".join(cells)
            identifiers = REQUIREMENT_ID.findall(joined)
            for identifier in identifiers:
                description = next(
                    (cell for cell in cells if cell and identifier not in cell),
                    joined.replace(identifier, "").strip(),
                )
                requirements.setdefault(identifier, description)

    for paragraph in document.paragraphs:
        text = clean(paragraph.text)
        for identifier in REQUIREMENT_ID.findall(text):
            requirements.setdefault(identifier, text.replace(identifier, "").strip())

    grouped = sorted(requirements.items(), key=lambda item: item[0])
    lines = [
        "# Requirements traceability matrix",
        "",
        f"Generated from `{SOURCE.name}`. Inventory contains **{len(grouped)} unique requirement IDs**.",
        "",
        "Status is deliberately conservative: `Partial` means at least one implementing control exists but acceptance evidence is incomplete; `Planned` means no completion claim is made.",
        "",
        "| Requirement | Summary | Status | Implementation evidence | Test evidence |",
        "|---|---|---|---|---|",
    ]

    for identifier, description in grouped:
        status, implementation, tests = evidence_for(identifier)
        lines.append(
            f"| {identifier} | {description or 'See authoritative SRS.'} | {status} | {implementation} | {tests} |"
        )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {len(grouped)} requirements to {OUTPUT}")


if __name__ == "__main__":
    main()
